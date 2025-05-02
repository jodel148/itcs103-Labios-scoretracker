import tkinter as tk
from openpyxl import load_workbook
from tkinter import messagebox

def validate_inputs():
    name = name_entry.get()
    score = score_entry.get()


    if not name or not score:
        messagebox.showerror("Input Error", "All fields are required!")
        return False

    else:
        return True


def save_file():
    if not validate_inputs:
        return
    

    name= name_entry.get()
    score= eval(score_entry.get())
    pas = "Passed"
    fail = "Failed"

    wb = load_workbook("Student_scores.xlsx")
    ws = wb["student_scores"]

    if score <= 74:
         ws.append([name, score, fail])
         wb.save("Student_scores.xlsx") 

    else:
        ws.append([name, score, pas])
        wb.save("Student_scores.xlsx") 
    
    messagebox.showinfo("Success", "Data saved successfully!")
    name_entry.delete(0, tk.END)
    score_entry.delete(0, tk.END)
   
def view_file():
   
    wl = load_workbook("Student_scores.xlsx")
    wb = wl.active

    data_window = tk.Toplevel(window)
    data_window.title("Student scores data")

    for i, row in enumerate(wb.iter_rows(values_only=True)):
        for j, value in enumerate(row):
            label = tk.Label(data_window, text=value, borderwidth=1, relief="solid", padx=6, pady=3)
            label.grid(row=i, column=j)



window = tk.Tk()
window.title("Score Tracker")

label1 = tk.Label(window, text="Name: ", font="Arial 12")
label2 = tk.Label(window, text="Score: ", font="Arial 12")

label1.grid(row=0, column=0)
label2.grid(row=1, column=0)

name_entry = tk.Entry(window)
score_entry = tk.Entry(window)

name_entry.grid(row=0, column=1)
score_entry.grid(row=1, column=1)

button1 = tk.Button(window, text="Submit", bg="yellow", fg="brown", command=save_file)
button2 = tk.Button(window, text="View", bg="yellow", fg="brown", command=view_file)

button1.grid(column=0, row=2, columnspan=2 )
button2.grid(column=0, row=3, columnspan=2 )







window.mainloop()


